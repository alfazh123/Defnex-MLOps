#!/usr/bin/env python3
"""Generate CSV and XLSX template files for the DEFNEX Dataset Intake Pipeline.

Run: python scripts/generate_templates.py
Output: templates/defnex_dataset_template.csv, templates/defnex_dataset_template.xlsx
"""

import csv
from pathlib import Path

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates"
COLUMNS = ["instruction", "input", "output", "style"]

SAMPLE_ROWS = [
    [
        "Jelaskan konsep LoRA",
        "",
        "LoRA (Low-Rank Adaptation) adalah metode parameter-efficient fine-tuning yang menyematkan matriks rank-rendah ke dalam layer transformer, memungkinkan fine-tuning dengan GPU minimum.",
        "formal",
    ],
    [
        "Apa itu SFT?",
        "dalam konteks machine learning",
        "Supervised Fine-Tuning (SFT) adalah proses melatih model bahasa besar menggunakan data berlabel (instruction-response pairs) untuk mengarahkan model agar menghasilkan output yang sesuai ekspektasi.",
        "formal",
    ],
    [
        "Buatkan kode Python untuk membaca file CSV",
        "",
        "import csv\n\nwith open('data.csv', 'r') as f:\n    reader = csv.DictReader(f)\n    for row in reader:\n        print(row)",
        "code",
    ],
    [
        "Apa perbedaan SFT dan RAG?",
        "untuk chatbot customer service",
        "SFT melatih model dengan data berlabel agar mengikuti format respons tertentu, sedangkan RAG mengambil dokumen relevan dari basis pengetahuan untuk menjawab pertanyaan tanpa melatih ulang model.",
        "formal",
    ],
]


def generate_csv() -> Path:
    csv_path = TEMPLATE_DIR / "defnex_dataset_template.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(COLUMNS)
        writer.writerows(SAMPLE_ROWS)
    return csv_path


def generate_xlsx() -> Path | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print(
            "⚠  openpyxl not installed, skipping XLSX generation. Run: pip install openpyxl"
        )
        return None

    xlsx_path = TEMPLATE_DIR / "defnex_dataset_template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Dataset Template"

    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    header_font = Font(bold=True, size=11, color="FFFFFF")

    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row_data in enumerate(SAMPLE_ROWS, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(max_len + 2, 60)

    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    return xlsx_path


def main():
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)

    csv_path = generate_csv()
    print(f"✓ CSV:  {csv_path}")

    xlsx_path = generate_xlsx()
    if xlsx_path:
        print(f"✓ XLSX: {xlsx_path}")

    print(f"\nColumns: {', '.join(COLUMNS)}")
    print(f"Sample rows: {len(SAMPLE_ROWS)}")


if __name__ == "__main__":
    main()
